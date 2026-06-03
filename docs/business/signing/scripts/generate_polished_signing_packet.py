from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[4]
OUT_DIR = ROOT / "docs/business/signing/ready-for-signature-2026-06-03"
LOGO = ROOT / "docs/business/brand-assets/images/mladis-connected-intelligence.png"

INK = "182131"
ACCENT = "1D5E7A"
ACCENT_DARK = "12384D"
MUTED = "64748B"
LIGHT = "EEF5F8"
LINE = "C7D7DF"
SOFT = "F7FAFC"


@dataclass(frozen=True)
class SignerBlock:
    label: str
    name: str
    role: str | None = None


@dataclass(frozen=True)
class DocSpec:
    source: str
    stem: str
    display_title: str
    subtitle: str
    doc_code: str
    parties: str
    prepared_for: str
    signers: tuple[SignerBlock, ...]


SPECS = (
    DocSpec(
        source="docs/business/operating-agreement/MLADIS_LLC_Operating_Agreement_DRAFT.md",
        stem="MLADIS_LLC_Operating_Agreement_READY_FOR_SIGNATURE",
        display_title="Operating Agreement",
        subtitle="MLADIS LLC",
        doc_code="MLADIS-GOV-001",
        parties="Piter Zacari Garcia Bautista, sole member",
        prepared_for="Internal company governance and bank/EIN readiness",
        signers=(SignerBlock("Member", "Piter Zacari Garcia Bautista", "Sole Member"),),
    ),
    DocSpec(
        source="docs/business/operating-agreement/initial-member-consent.md",
        stem="MLADIS_LLC_Initial_Member_Consent_READY_FOR_SIGNATURE",
        display_title="Initial Written Consent",
        subtitle="Sole Member Resolutions",
        doc_code="MLADIS-GOV-002",
        parties="Piter Zacari Garcia Bautista, sole member",
        prepared_for="Formation, operating authority, EIN, banking, systems, publication, and records",
        signers=(SignerBlock("Member", "Piter Zacari Garcia Bautista", "Sole Member"),),
    ),
    DocSpec(
        source="docs/business/operations/mladis-bookings-founding-operations-pillar-acknowledgment.md",
        stem="MLADIS_Bookings_Founding_Operations_Pillar_Acknowledgment_READY_FOR_SIGNATURE",
        display_title="Founding Operations Pillar Acknowledgment",
        subtitle="MLADIS Bookings Services Understanding",
        doc_code="MLADIS-BKG-001",
        parties="Piter Zacari Garcia Bautista, MLADIS LLC, and Diana Sori Garcia Bautista",
        prepared_for="Diana signature and Dropbox Sign acceptance",
        signers=(
            SignerBlock("Piter Zacari Garcia Bautista", "Piter Zacari Garcia Bautista", "Individually"),
            SignerBlock("MLADIS LLC", "Piter Zacari Garcia Bautista", "Sole Member, MLADIS LLC"),
            SignerBlock("Diana Sori Garcia Bautista", "Diana Sori Garcia Bautista", "Dominican Republic Operations Lead"),
        ),
    ),
)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LINE, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_width(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def add_bottom_border(paragraph, color: str = ACCENT, size: str = "18") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)


def set_run(run, *, size: int | float | None = None, color: str | None = None, bold: bool | None = None):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in [
        ("Heading 1", 15, ACCENT_DARK, 14, 6),
        ("Heading 2", 12.5, ACCENT, 10, 5),
        ("Heading 3", 11.5, ACCENT_DARK, 8, 4),
    ]:
        s = styles[style_name]
        s.font.name = "Aptos Display"
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
        s.font.size = Pt(size)
        s.font.color.rgb = RGBColor.from_string(color)
        s.font.bold = True
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True


def add_header_footer(doc: Document, spec: DocSpec) -> None:
    for section in doc.sections:
        header = section.header
        if header.paragraphs:
            p = header.paragraphs[0]
        else:
            p = header.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run("MLADIS LLC  |  ")
        set_run(r, size=8.5, color=MUTED, bold=True)
        r = p.add_run(spec.doc_code)
        set_run(r, size=8.5, color=MUTED)

        footer = section.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("Private company record - ready for signature")
        set_run(r, size=8, color=MUTED)


def add_cover(doc: Document, spec: DocSpec) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if LOGO.exists():
        p.add_run().add_picture(str(LOGO), width=Inches(0.82))

    brand = doc.add_paragraph()
    brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = brand.add_run("MLADIS LLC")
    set_run(r, size=12, color=ACCENT_DARK, bold=True)
    brand.paragraph_format.space_after = Pt(2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run(spec.display_title)
    set_run(r, size=25, color=INK, bold=True)
    title.paragraph_format.space_after = Pt(2)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run(spec.subtitle)
    set_run(r, size=13, color=ACCENT)
    subtitle.paragraph_format.space_after = Pt(18)
    add_bottom_border(subtitle, ACCENT, "12")

    rows = [
        ("Document code", spec.doc_code),
        ("Prepared for", spec.prepared_for),
        ("Parties / signer context", spec.parties),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_width(table, [2250, 6500])
    for row, (label, value) in zip(table.rows, rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
        set_cell_shading(row.cells[0], LIGHT)
        p = row.cells[0].paragraphs[0]
        r = p.add_run(label.upper())
        set_run(r, size=8.5, color=ACCENT_DARK, bold=True)
        p = row.cells[1].paragraphs[0]
        r = p.add_run(value)
        set_run(r, size=10.5, color=INK)

    doc.add_paragraph()
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = note.add_run("Prepared for execution through the approved MLADIS e-signature workflow.")
    set_run(r, size=9.5, color=MUTED, bold=True)
    note.paragraph_format.space_before = Pt(6)

    doc.add_section(WD_SECTION.NEW_PAGE)


def parse_source(path: Path) -> list[tuple[str, str | list[str]]]:
    lines = path.read_text(encoding="utf-8").splitlines()
    blocks: list[tuple[str, str | list[str]]] = []
    table: list[str] = []

    def flush_table():
        nonlocal table
        if table:
            blocks.append(("table", table))
            table = []

    for raw in lines:
        line = raw.strip()
        if not line:
            flush_table()
            continue
        if line.startswith("Draft for review"):
            continue
        if line.startswith("Not legal"):
            continue
        if line.startswith("Do not sign"):
            continue
        line = line.replace("as of ____________, 2026", "as of June 3, 2026")
        line = line.replace("dated ____________, 2026", "dated June 3, 2026")
        line = line.replace("entered into as of ____________, 2026", "entered into as of June 3, 2026")
        line = line.replace("This Operating Agreement is entered into as of ____________, 2026", "This Operating Agreement is entered into as of June 3, 2026")
        if line.startswith("# "):
            continue
        if line.startswith("## ") and "signature" in line.lower():
            flush_table()
            break
        if line.startswith("## "):
            flush_table()
            blocks.append(("h1", line[3:].strip()))
        elif line.startswith("### "):
            flush_table()
            blocks.append(("h2", line[4:].strip()))
        elif line.startswith("|"):
            table.append(line)
        elif line.startswith("- "):
            flush_table()
            blocks.append(("bullet", line[2:].strip()))
        elif re.match(r"^\d+\.\s+", line):
            flush_table()
            blocks.append(("number", re.sub(r"^\d+\.\s+", "", line).strip()))
        else:
            flush_table()
            blocks.append(("p", line))
    flush_table()
    return blocks


def add_note_box(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [9000])
    cell = table.cell(0, 0)
    set_cell_shading(cell, SOFT)
    set_cell_border(cell, LINE)
    p = cell.paragraphs[0]
    r = p.add_run(text)
    set_run(r, size=9.5, color=MUTED, bold=True)
    p.paragraph_format.space_after = Pt(0)


def add_markdown_table(doc: Document, lines: list[str]) -> None:
    rows = []
    for line in lines:
        cells = [c.strip() for c in line.strip("|").split("|")]
        if all(set(c) <= {"-", ":", " "} for c in cells):
            continue
        rows.append(cells)
    if not rows:
        return
    max_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=max_cols)
    width = 9000
    widths = [int(width / max_cols)] * max_cols
    set_table_width(table, widths)
    for ridx, row_data in enumerate(rows):
        for cidx in range(max_cols):
            cell = table.cell(ridx, cidx)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx == 0:
                set_cell_shading(cell, LIGHT)
            text = row_data[cidx] if cidx < len(row_data) else ""
            p = cell.paragraphs[0]
            r = p.add_run(text)
            set_run(r, size=9.2 if ridx else 8.7, color=ACCENT_DARK if ridx == 0 else INK, bold=(ridx == 0))
            p.paragraph_format.space_after = Pt(0)
    doc.add_paragraph()


def add_body(doc: Document, spec: DocSpec) -> None:
    for kind, value in parse_source(ROOT / spec.source):
        text = value if isinstance(value, str) else ""
        if kind == "h1":
            doc.add_paragraph(text, style="Heading 1")
        elif kind == "h2":
            doc.add_paragraph(text, style="Heading 2")
        elif kind == "p":
            p = doc.add_paragraph()
            r = p.add_run(text)
            set_run(r, size=10.5, color=INK)
        elif kind == "bullet":
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(text)
            set_run(r, size=10.2, color=INK)
        elif kind == "number":
            p = doc.add_paragraph(style="List Number")
            r = p.add_run(text)
            set_run(r, size=10.2, color=INK)
        elif kind == "table" and isinstance(value, list):
            add_markdown_table(doc, value)


def add_signature_block(doc: Document, spec: DocSpec) -> None:
    for idx, signer in enumerate(spec.signers):
        if idx == 0:
            doc.add_page_break()
        else:
            doc.add_page_break()
        doc.add_paragraph("Signature Authorization", style="Heading 1")
        p = doc.add_paragraph()
        r = p.add_run(
            f"{signer.label} signs below to acknowledge and execute this document as routed through the approved e-signature workflow."
        )
        set_run(r, size=10.5, color=INK)
        table = doc.add_table(rows=5, cols=2)
        set_table_width(table, [2250, 6500])
        fields = [
            ("Signer", signer.label),
            ("Name", signer.name),
            ("Role", signer.role or ""),
            ("Signature", ""),
            ("Date", ""),
        ]
        for row, (label, value) in zip(table.rows, fields):
            prevent_row_split(row)
            for cell in row.cells:
                set_cell_border(cell)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(row.cells[0], LIGHT)
            lp = row.cells[0].paragraphs[0]
            lr = lp.add_run(label.upper())
            set_run(lr, size=8.5, color=ACCENT_DARK, bold=True)
            vp = row.cells[1].paragraphs[0]
            if value:
                vr = vp.add_run(value)
                set_run(vr, size=10.2, color=INK, bold=label == "Signer")
            else:
                vr = vp.add_run(" ")
                set_run(vr, size=18, color=INK)
            vp.paragraph_format.space_after = Pt(0)
        doc.add_paragraph()


def build_doc(spec: DocSpec) -> Path:
    doc = Document()
    style_document(doc)
    add_header_footer(doc, spec)
    add_cover(doc, spec)
    add_body(doc, spec)
    add_signature_block(doc, spec)

    doc.core_properties.author = "MLADIS LLC"
    doc.core_properties.title = f"{spec.subtitle} - {spec.display_title}"
    doc.core_properties.subject = spec.doc_code
    doc.core_properties.keywords = "MLADIS LLC; signing packet; Dropbox Sign"

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUT_DIR / f"{spec.stem}.docx"
    doc.save(out)
    return out


def write_readme(outputs: list[Path]) -> None:
    rows = []
    for path in outputs:
        rows.append(f"| `{path.name}` | `{sha256(path)}` |")
    readme = f"""# Ready For Signature Packet - 2026-06-03

Private ready-for-signature packet for MLADIS LLC.

## Status

- Prepared on: 2026-06-03
- Visual system: branded MLADIS legal/business packet
- Signature images: none embedded
- Intended signing tool: Dropbox Sign under `garcp37@mladis.com`
- Status: ready for signature field placement and sending

## Files

| File | SHA-256 |
| --- | --- |
{chr(10).join(rows)}

## Signing Instructions

- Piter signs the Operating Agreement and Initial Written Consent.
- Piter signs the Diana Acknowledgment individually and for MLADIS LLC.
- Diana signs only the Diana Acknowledgment.
- Do not reuse, crop, or paste Diana's prior signature from archived source documents. Let Diana sign through Dropbox Sign.

## After Signatures

Save the completed signed PDFs and Dropbox Sign audit trail in private company records. Do not commit reusable signature images, identity documents, EIN letters, SSNs, bank records, tax returns, or payment credentials.
"""
    (OUT_DIR / "README.md").write_text(readme, encoding="utf-8")


def main() -> None:
    outputs = [build_doc(spec) for spec in SPECS]
    write_readme(outputs)
    for path in outputs:
        print(path)


if __name__ == "__main__":
    main()
